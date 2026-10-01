import streamlit as st
import os
import time
from docx import Document
from io import BytesIO
from main import run_crew

# ==========================================
# PAGE CONFIGURATION & UI
# ==========================================
st.set_page_config(
    page_title="FreelanceOS | Agentic Copilot",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for modern UI
st.markdown("""
<style>
    .main-header {
        background: linear-gradient(90deg, #6366F1 0%, #A855F7 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-size: 3rem;
        font-weight: 800;
        margin-bottom: 0px;
    }
    .sub-header { color: #64748b; font-size: 1.2rem; margin-bottom: 2rem; }
    .metric-card {
        background-color: #ffffff;
        border-radius: 10px;
        padding: 20px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
        text-align: center;
        border: 1px solid #e2e8f0;
    }
    .metric-value { font-size: 2rem; font-weight: bold; color: #0F172A; }
    .metric-label { font-size: 0.9rem; color: #64748b; text-transform: uppercase; }
    
    @media (prefers-color-scheme: dark) {
        .metric-card { background-color: #1E293B; border-color: #334155; }
        .metric-value { color: #F8FAFC; }
        .metric-label { color: #94A3B8; }
    }
</style>
""", unsafe_allow_html=True)

# ==========================================
# SECURITY & STATE MANAGEMENT
# ==========================================
def get_server_api_key():
    try:
        return st.secrets.get("GEMINI_API_KEY")
    except Exception:
        return os.environ.get("GEMINI_API_KEY", "")

if "run_count" not in st.session_state:
    st.session_state.run_count = 0
if "last_run_time" not in st.session_state:
    st.session_state.last_run_time = 0
if "results" not in st.session_state:
    st.session_state.results = None
if "job_input" not in st.session_state:
    st.session_state.job_input = ""

# ==========================================
# SIDEBAR
# ==========================================
with st.sidebar:
    st.image("https://img.icons8.com/fluency/96/robot-2.png", width=60)
    st.markdown("### FreelanceOS")
    st.caption("Your Agentic Copilot")
    
    st.divider()
    st.markdown("#### How it works")
    st.markdown("""
    1. **Lead Scout** analyzes feasibility.
    2. **Proposal Architect** drafts a pitch.
    3. **Project Manager** plans milestones.
    4. **Finance Officer** prices the job.
    """)
    st.divider()
    
    st.markdown("#### Settings")
    user_api_key = st.text_input(
        "Use your own Gemini API Key (Optional)", 
        type="password", 
        help="Leave empty to use the free public quota. Key is not stored."
    )
    
    tone_selector = st.selectbox("Proposal Tone", ["professional", "friendly", "bold"])
    
    st.divider()
    st.metric("Session Runs", f"{st.session_state.run_count} / 5")

# ==========================================
# MAIN LAYOUT
# ==========================================
st.markdown('<div class="main-header">🤖 FreelanceOS</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-header">Instantly analyze jobs, generate proposals, and plan projects using AI agents.</div>', unsafe_allow_html=True)

st.markdown("### 📋 Job & Profile Details")

job_description = st.text_area(
    "Job posting", 
    value=st.session_state.job_input, 
    height=150, 
    placeholder="Paste the full job description here...",
    max_chars=4000
)

# New fields from the provided UI image
col1, col2 = st.columns(2)
with col1:
    freelancer_name = st.text_input("Your name", value="Alex")
    freelancer_skills = st.text_input("Your skills", value="Python, React, copywriting")
    freelancer_experience = st.text_input("Experience", value="3 years freelancing")
with col2:
    freelancer_rate = st.number_input("Hourly rate (USD)", value=50.00, step=1.0)
    freelancer_availability = st.text_input("Availability", value="20 hours per week")

st.write("")
if st.button("🚀 Run Agentic Copilot", type="primary", use_container_width=True):
    if len(job_description) < 10:
        st.warning("Please enter a valid job description.")
        st.stop()
        
    current_time = time.time()
    if current_time - st.session_state.last_run_time < 60:
        st.error(f"Cooldown active. Please wait {int(60 - (current_time - st.session_state.last_run_time))} seconds.")
        st.stop()
        
    if st.session_state.run_count >= 5 and not user_api_key:
        st.error("Max runs reached for this free session. Please provide your own Gemini API key in the sidebar to continue.")
        st.stop()

    active_key = user_api_key.strip() if user_api_key.strip() else get_server_api_key()
    
    if not active_key:
        st.error("Server API key not configured and no custom key provided. Please check secrets.")
        st.stop()

    with st.status("🧠 Agents are analyzing the job...", expanded=True) as status:
        st.write("🕵️‍♂️ **Lead Scout** is analyzing fit & red flags...")
        st.write("✍️️ **Proposal Architect** is drafting the pitch...")
        st.write("📊 **Project Manager** is creating milestones...")
        st.write("💰 **Finance Officer** is crunching numbers...")
        
        # Pass all the new fields to the runner
        profile_data = {
            "name": freelancer_name,
            "rate": freelancer_rate,
            "skills": freelancer_skills,
            "availability": freelancer_availability,
            "experience": freelancer_experience,
            "tone": tone_selector
        }
        
        success, msg, results = run_crew(job_description, active_key, profile_data)
        
        if success:
            status.update(label="✅ Analysis Complete!", state="complete", expanded=False)
            st.session_state.results = results
            st.session_state.run_count += 1
            st.session_state.last_run_time = current_time
        else:
            status.update(label="❌ Error occurred", state="error", expanded=True)
            st.error(msg)

# ==========================================
# RESULTS DASHBOARD
# ==========================================
if st.session_state.results:
    results = st.session_state.results
    metrics = results.get("metrics", {})
    
    st.markdown("### 📊 Dashboard Overview")
    m1, m2, m3, m4 = st.columns(4)
    
    with m1:
        st.markdown(f"""<div class="metric-card">
            <div class="metric-label">Fit Score</div>
            <div class="metric-value">{metrics.get('fit_score_1_to_10', 'N/A')}/10</div>
            </div>""", unsafe_allow_html=True)
    with m2:
        decision = metrics.get('decision', 'N/A')
        color = "#10B981" if "GO" in decision.upper() and "NO" not in decision.upper() else "#EF4444"
        st.markdown(f"""<div class="metric-card">
            <div class="metric-label">Decision</div>
            <div class="metric-value" style="color: {color};">{decision}</div>
            </div>""", unsafe_allow_html=True)
    with m3:
        st.markdown(f"""<div class="metric-card">
            <div class="metric-label">Est. Hours</div>
            <div class="metric-value">{metrics.get('estimated_hours', 'N/A')}</div>
            </div>""", unsafe_allow_html=True)
    with m4:
        st.markdown(f"""<div class="metric-card">
            <div class="metric-label">Suggested Price</div>
            <div class="metric-value">${metrics.get('recommended_price_usd', 'N/A')}</div>
            </div>""", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    
    t1, t2, t3, t4 = st.tabs(["🕵️ Lead Analysis", "📝 Proposal", "📊 Project Plan", "💰 Finances"])
    
    with t1:
        st.markdown(results["lead_analysis"])
    
    with t2:
        st.markdown("#### Copy-ready Proposal")
        st.info("Review and adjust before sending to the client.")
        st.code(results["proposal"], language="markdown")

    with t3:
        st.markdown(results["project_plan"])
        
    with t4:
        st.markdown(results["finance"])

    st.divider()
    st.markdown("### 💾 Export Reports")
    
    full_md = f"""# FreelanceOS Report\n## 1. Analysis\n{results['lead_analysis']}\n## 2. Proposal\n{results['proposal']}\n## 3. Project Plan\n{results['project_plan']}\n## 4. Finance & Terms\n{results['finance']}"""
    
    def create_docx(markdown_text):
        doc = Document()
        doc.add_heading('FreelanceOS Agentic Report', 0)
        for line in markdown_text.split('\n'):
            if line.startswith('## '):
                doc.add_heading(line.replace('## ', ''), level=1)
            elif line.startswith('### '):
                doc.add_heading(line.replace('### ', ''), level=2)
            elif line.strip():
                doc.add_paragraph(line)
        bio = BytesIO()
        doc.save(bio)
        return bio.getvalue()

    c1, c2, c3 = st.columns([1,1,2])
    with c1:
        st.download_button("⬇️ Download Markdown", full_md, file_name="freelanceos_report.md", mime="text/markdown")
    with c2:
        st.download_button("⬇️ Download DOCX", create_docx(full_md), file_name="freelanceos_report.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")